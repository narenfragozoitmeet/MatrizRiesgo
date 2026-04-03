import pdfplumber
from docx import Document
import openpyxl
from io import BytesIO
import logging

logger = logging.getLogger(__name__)

def parse_pdf(data: bytes) -> str:
    """Extract text from PDF"""
    try:
        text_parts = []
        with pdfplumber.open(BytesIO(data)) as pdf:
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    text_parts.append(text)
        return "\n".join(text_parts)
    except Exception as e:
        logger.error(f"Error parsing PDF: {e}")
        raise

def parse_docx(data: bytes) -> str:
    """Extract text from Word document"""
    try:
        doc = Document(BytesIO(data))
        text_parts = []
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)
        return "\n".join(text_parts)
    except Exception as e:
        logger.error(f"Error parsing DOCX: {e}")
        raise

def parse_xlsx(data: bytes) -> str:
    """Extract text from Excel file"""
    try:
        wb = openpyxl.load_workbook(BytesIO(data), read_only=True)
        text_parts = []
        for sheet in wb.worksheets:
            text_parts.append(f"\n=== Hoja: {sheet.title} ===\n")
            for row in sheet.iter_rows(values_only=True):
                row_text = "\t".join([str(cell) if cell is not None else "" for cell in row])
                if row_text.strip():
                    text_parts.append(row_text)
        return "\n".join(text_parts)
    except Exception as e:
        logger.error(f"Error parsing XLSX: {e}")
        raise

def parse_document(data: bytes, content_type: str, filename: str) -> str:
    """Parse document based on content type"""
    if content_type == 'application/pdf':
        return parse_pdf(data)
    elif content_type in ['application/vnd.openxmlformats-officedocument.wordprocessingml.document', 'application/msword']:
        return parse_docx(data)
    elif content_type in ['application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'application/vnd.ms-excel']:
        return parse_xlsx(data)
    else:
        raise ValueError(f"Unsupported document type: {content_type}")